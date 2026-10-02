"""Report Service — deterministic report generation & export for RIVA AI.

Generates financial reports (P&L, Cash Flow, Expense, Transaction,
Balance Sheet, Trial Balance) from the accounting engine's real data
and exports them in common formats (CSV, XLSX, PDF).

All calculations are deterministic and come from analytics_service.
The AI can request report generation through the report tools; the
actual files are produced here, never faked.

Report formats:
- CSV: stdlib csv
- XLSX: openpyxl
- PDF: reportlab
"""

import csv
import io
from datetime import datetime
from decimal import Decimal
from typing import BinaryIO

from sqlalchemy.orm import Session

from analytics_service import (
    get_period_totals,
    get_expense_breakdown,
    get_transactions,
    get_account_balances,
    AccountBalance,
)
from financial_periods import Period


# ---------------------------------------------------------------------------
# Report definitions (which reports exist)
# ---------------------------------------------------------------------------

# Reports we can generate with real data. `key` is used by AI tools and
# the export endpoint.
AVAILABLE_REPORTS = {
    "profit_loss": {
        "label": "Profit & Loss",
        "label_fa": "صورت سود و زیان",
        "description": "Revenue, expenses and net income for a period.",
    },
    "cash_flow": {
        "label": "Cash Flow",
        "label_fa": "جریان نقدی",
        "description": "Net cash generated/(used) for a period.",
    },
    "expense_report": {
        "label": "Expense Report",
        "label_fa": "گزارش هزینه‌ها",
        "description": "Expenses grouped by category with amounts and percentages.",
    },
    "transaction_report": {
        "label": "Transaction Report",
        "label_fa": "گزارش تراکنش‌ها",
        "description": "List of transactions in a period.",
    },
    "balance_sheet": {
        "label": "Balance Sheet",
        "label_fa": "ترازنامه",
        "description": "Assets, liabilities and equity at a point in time.",
    },
    "trial_balance": {
        "label": "Trial Balance",
        "label_fa": "تراز آزمایشی",
        "description": "All account balances with debit/credit classification.",
    },
    "account_statement": {
        "label": "Account Statement",
        "label_fa": "صورت حساب",
        "description": "Transactions in a specific account.",
    },
}

REPORT_KEYS = set(AVAILABLE_REPORTS.keys())


def resolve_report_key(name: str) -> str | None:
    """Resolve a user/AI supplied report name to a canonical report key."""
    normalized = name.strip().lower().replace("_", " ").replace("-", " ")
    normalized = " ".join(normalized.split())

    # Direct match on key
    for key in REPORT_KEYS:
        if normalized == key.replace("_", " "):
            return key

    # Alias matching
    aliases = {
        "profit_loss": ["profit and loss", "pnl", "p&l", "income statement", "سود و زیان"],
        "cash_flow": ["cashflow", "statement of cash flows", "جریان نقدی"],
        "expense_report": ["expenses", "expense", "هزینه", "هزینه ها"],
        "transaction_report": ["transactions", "transaction", "تراکنش"],
        "balance_sheet": ["balance", "ترازنامه"],
        "trial_balance": ["trial balance", "تراز آزمایشی"],
        "account_statement": ["account statement", "صورت حساب"],
    }
    for key, alias_list in aliases.items():
        if any(a in name.lower() for a in alias_list):
            return key

    return None


# ---------------------------------------------------------------------------
# Structured report data
# ---------------------------------------------------------------------------

def build_report_data(
    db: Session,
    report_key: str,
    business_id: str,
    business_name: str,
    base_currency: str,
    period: Period,
    account_name: str | None = None,
) -> dict:
    """Build structured, deterministic report data for a report key.

    Returns a dict that serializes to JSON and can be rendered as a table
    or exported to CSV/XLSX/PDF.
    """
    totals = get_period_totals(db, business_id, period)
    balances = get_account_balances(db, business_id)

    if report_key == "profit_loss":
        return _profit_loss(business_name, base_currency, period, totals)
    if report_key == "cash_flow":
        return _cash_flow(business_name, base_currency, period, totals)
    if report_key == "expense_report":
        expenses = get_expense_breakdown(db, business_id, period)
        return _expense_report(business_name, base_currency, period, totals, expenses)
    if report_key == "transaction_report":
        transactions = get_transactions(db, business_id, period, limit=200)
        return _transaction_report(business_name, base_currency, period, transactions)
    if report_key == "balance_sheet":
        return _balance_sheet(business_name, base_currency, period, balances)
    if report_key == "trial_balance":
        return _trial_balance(business_name, base_currency, period, balances)
    if report_key == "account_statement":
        transactions = get_transactions(db, business_id, period, category=account_name, limit=200)
        return _transaction_report(business_name, base_currency, period, transactions, account_name=account_name)
    raise ValueError(f"Unknown report: {report_key}")


def _profit_loss(business_name, base_currency, period, totals) -> dict:
    net = totals.revenue - totals.expenses
    return {
        "title": "Profit & Loss",
        "business": business_name,
        "currency": base_currency,
        "period": period.label,
        "generated_at": datetime.utcnow().isoformat(),
        "rows": [
            {"label": "Revenue", "amount": str(totals.revenue)},
            {"label": "Expenses", "amount": str(totals.expenses)},
            {"label": "Net Income", "amount": str(net)},
        ],
        "summary": {
            "revenue": str(totals.revenue),
            "expenses": str(totals.expenses),
            "net_income": str(net),
            "transaction_count": totals.transaction_count,
        },
    }


def _cash_flow(business_name, base_currency, period, totals) -> dict:
    return {
        "title": "Cash Flow",
        "business": business_name,
        "currency": base_currency,
        "period": period.label,
        "generated_at": datetime.utcnow().isoformat(),
        "rows": [
            {"label": "Operating Cash Flow (revenue − expenses)", "amount": str(totals.cash_flow)},
        ],
        "summary": {
            "cash_flow": str(totals.cash_flow),
            "revenue": str(totals.revenue),
            "expenses": str(totals.expenses),
            "transaction_count": totals.transaction_count,
        },
    }


def _expense_report(business_name, base_currency, period, totals, expenses) -> dict:
    rows = [
        {"category": e.category, "amount": str(e.amount), "percentage": f"{e.percentage}%"}
        for e in expenses
    ]
    return {
        "title": "Expense Report",
        "business": business_name,
        "currency": base_currency,
        "period": period.label,
        "generated_at": datetime.utcnow().isoformat(),
        "rows": rows,
        "summary": {
            "total_expenses": str(totals.expenses),
            "category_count": len(rows),
        },
    }


def _transaction_report(business_name, base_currency, period, transactions, account_name=None) -> dict:
    rows = [
        {
            "date": t.date.isoformat(),
            "category": t.category,
            "description": t.description,
            "amount": str(t.amount),
        }
        for t in transactions
    ]
    return {
        "title": f"Transaction Report{f' — {account_name}' if account_name else ''}",
        "business": business_name,
        "currency": base_currency,
        "period": period.label,
        "generated_at": datetime.utcnow().isoformat(),
        "rows": rows,
        "summary": {
            "transaction_count": len(rows),
        },
    }


def _balance_sheet(business_name, base_currency, period, balances) -> dict:
    assets = [b for b in balances if b.account_type == "asset"]
    liabilities = [b for b in balances if b.account_type == "liability"]
    equity = [b for b in balances if b.account_type == "equity"]

    def total(list_):
        return sum((b.balance for b in list_), Decimal("0"))

    assets_total = total(assets)
    liabilities_total = total(liabilities)
    equity_total = total(equity)

    rows = []
    for b in assets:
        rows.append({"section": "Assets", "account": b.name, "amount": str(b.balance)})
    for b in liabilities:
        rows.append({"section": "Liabilities", "account": b.name, "amount": str(b.balance)})
    for b in equity:
        rows.append({"section": "Equity", "account": b.name, "amount": str(b.balance)})

    return {
        "title": "Balance Sheet",
        "business": business_name,
        "currency": base_currency,
        "period": "as of now",
        "generated_at": datetime.utcnow().isoformat(),
        "rows": rows,
        "summary": {
            "assets": str(assets_total),
            "liabilities": str(liabilities_total),
            "equity": str(equity_total),
        },
    }


def _trial_balance(business_name, base_currency, period, balances) -> dict:
    rows = []
    for b in balances:
        if _debit_normal(b.account_type):
            debit = str(b.balance) if b.balance > 0 else "0"
            credit = str(-b.balance) if b.balance < 0 else "0"
        else:
            credit = str(b.balance) if b.balance > 0 else "0"
            debit = str(-b.balance) if b.balance < 0 else "0"
        rows.append({
            "account": b.name,
            "type": b.account_type,
            "debit": debit,
            "credit": credit,
        })
    return {
        "title": "Trial Balance",
        "business": business_name,
        "currency": base_currency,
        "period": "as of now",
        "generated_at": datetime.utcnow().isoformat(),
        "rows": rows,
        "summary": {"account_count": len(rows)},
    }


def _debit_normal(account_type: str) -> bool:
    return account_type in {"asset", "expense"}


# ---------------------------------------------------------------------------
# Export serialization — CSV / XLSX / PDF
# ---------------------------------------------------------------------------

def _flatten_rows(report: dict) -> list[dict]:
    """Flatten report data into exportable row dicts."""
    summary = report.get("summary", {})
    if report.get("rows"):
        return report["rows"]

    # No rows — emit summary as a single row
    return [summary]


def _csv_bytes(report: dict) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow([report["title"]])
    writer.writerow(["Business", report.get("business", "")])
    writer.writerow(["Currency", report.get("currency", "")])
    writer.writerow(["Period", report.get("period", "")])
    writer.writerow([""])

    rows = _flatten_rows(report)
    if rows:
        headers = list(rows[0].keys())
        writer.writerow(headers)
        for row in rows:
            writer.writerow([row.get(h, "") for h in headers])

    # Summary
    summary = report.get("summary", {})
    if summary:
        writer.writerow([""])
        writer.writerow(["Summary"])
        for k, v in summary.items():
            writer.writerow([k, v])

    return buffer.getvalue().encode("utf-8")


def _xlsx_bytes(report: dict) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        raise ValueError(
            "Excel export is not available — the 'openpyxl' package is not installed on the server."
        )

    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:30]

    bold = Font(bold=True)

    ws.append([report["title"]])
    ws.append(["Business", report.get("business", "")])
    ws.append(["Currency", report.get("currency", "")])
    ws.append(["Period", report.get("period", "")])
    ws.append([])

    rows = _flatten_rows(report)
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=c).font = bold
        for row in rows:
            ws.append([row.get(h, "") for h in headers])

    summary = report.get("summary", {})
    if summary:
        ws.append([])
        ws.append(["Summary"])
        ws.cell(row=ws.max_row, column=1).font = bold
        for k, v in summary.items():
            ws.append([k, v])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(report: dict) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        from reportlab.lib import colors
    except ImportError:
        raise ValueError(
            "PDF export is not available — the 'reportlab' package is not installed on the server."
        )

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    margin = 20 * mm
    y = height - margin

    c.setFont("Helvetica-Bold", 16)
    c.drawString(margin, y, report["title"])
    y -= 8 * mm

    c.setFont("Helvetica", 10)
    c.drawString(margin, y, f"Business: {report.get('business', '')}")
    y -= 5 * mm
    c.drawString(margin, y, f"Currency: {report.get('currency', '')}  |  Period: {report.get('period', '')}")
    y -= 5 * mm
    c.drawString(margin, y, f"Generated: {report.get('generated_at', '')}")
    y -= 10 * mm

    rows = _flatten_rows(report)
    if not rows:
        c.drawString(margin, y, "No data for this report.")
        c.save()
        return buffer.getvalue()

    headers = list(rows[0].keys())
    col_width = (width - 2 * margin) / len(headers)

    # Header
    c.setFont("Helvetica-Bold", 8)
    x = margin
    for h in headers:
        c.drawString(x, y, str(h)[:24])
        x += col_width
    y -= 4 * mm
    c.setStrokeColor(colors.grey)
    c.setLineWidth(0.5)
    c.line(margin, y, width - margin, y)
    y -= 4 * mm

    # Rows
    c.setFont("Helvetica", 8)
    for row in rows:
        if y < margin:
            c.showPage()
            c.setFont("Helvetica-Bold", 8)
            y = height - margin
            x = margin
            for h in headers:
                c.drawString(x, y, str(h)[:24])
                x += col_width
            y -= 6 * mm
            c.setFont("Helvetica", 8)

        x = margin
        for h in headers:
            val = str(row.get(h, ""))
            c.drawString(x, y, val[:24])
            x += col_width
        y -= 4 * mm

    # Summary
    summary = report.get("summary", {})
    if summary:
        y -= 6 * mm
        c.setFont("Helvetica-Bold", 10)
        c.drawString(margin, y, "Summary")
        y -= 5 * mm
        c.setFont("Helvetica", 9)
        for k, v in summary.items():
            c.drawString(margin, y, f"{k}: {v}")
            y -= 4 * mm

    c.save()
    return buffer.getvalue()


def export_report(
    report: dict,
    fmt: str,
) -> tuple[bytes, str]:
    """Serialize a report dict into bytes for the requested format.

    Args:
        report: Structured report dict from build_report_data.
        fmt: 'csv', 'xlsx', or 'pdf'.

    Returns:
        (bytes, content_type)
    """
    fmt = fmt.lower()
    if fmt == "csv":
        return _csv_bytes(report), "text/csv"
    if fmt == "xlsx":
        return _xlsx_bytes(report), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "pdf":
        return _pdf_bytes(report), "application/pdf"
    raise ValueError(f"Unsupported format: {fmt}")


def default_filename(report: dict, fmt: str) -> str:
    """Build a safe filename for a report export."""
    stamp = datetime.utcnow().strftime("%Y-%m-%d")
    safe_title = report.get("title", "report").lower().replace(" ", "-")
    safe_title = "".join(c for c in safe_title if c.isalnum() or c in "-_")
    return f"riva-{safe_title}-{stamp}.{fmt.lower()}"